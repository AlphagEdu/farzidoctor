from flask import Flask, request, send_file
from docx import Document
import io

app = Flask(__name__)

# Load the template document
TEMPLATE_PATH = "report_draft.docx"  # Change to your actual file name

@app.route("/generate", methods=["POST"])
def generate_document():
    data = request.json  # Get JSON data from frontend
    doc = Document(TEMPLATE_PATH)  # Load the Word document

    # Replace placeholders with actual values
    replacements = {
        "NAME_PLACEHOLDER": data["name"],
        "REGN_DATE_PLACEHOLDER": data["regn_date"],
        "SAMPLE_COLLECTION_PLACEHOLDER": data["sample_collection"],
        "PRINT_DATE_PLACEHOLDER": data["print_date"],
        "AGE_SEX_PLACEHOLDER": f"{data['age']} Years / {data['gender']}",
    }

    for para in doc.paragraphs:
        for key, value in replacements.items():
            if key in para.text:
                para.text = para.text.replace(key, value)

    # Save to a buffer instead of disk
    file_stream = io.BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)

    return send_file(file_stream, as_attachment=True, download_name="Updated_Report.docx", mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document")

if __name__ == "__main__":
    app.run(debug=True)
